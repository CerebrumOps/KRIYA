import openpyxl

path = '/home/vicky/Projects/KRIYA/workspace/CDU-II_C101_Overhead_Vapor_Line_Integrity_Cost_Workbook.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb['Cost & Severity Breakdown']

# Clear old data rows (4,5,6)
for r in [6, 5, 4]:
    for c in range(1, 8):
        ws.cell(row=r, column=c).value = None

# Write the single in-scope C-101 line item
ws.cell(row=4, column=1, value='C-101')
ws.cell(row=4, column=2, value='CDU-II')
ws.cell(row=4, column=3, value='Tray 44 Nozzle Spool (Salt Weeping / Category-A Defect)')
ws.cell(row=4, column=4, value='CRITICAL')
ws.cell(row=4, column=5, value=180000)   # Labor
ws.cell(row=4, column=6, value=460000)  # Materials
ws.cell(row=4, column=7, value='=E4+F4') # Total formula

# Clear the TOTAL row (row 7)
for c in range(1, 8):
    ws.cell(row=7, column=c).value = None

# Add a total row at row 5
ws.cell(row=5, column=1, value='TOTAL REMEDIATION EXPENDITURE')
ws.cell(row=5, column=5, value='=E4')
ws.cell(row=5, column=6, value='=F4')
ws.cell(row=5, column=7, value='=G4')

wb.save(path)
print('Workbook rewritten with C-101-only scope.')

# Verify
wb2 = openpyxl.load_workbook(path, data_only=False)
ws2 = wb2['Cost & Severity Breakdown']
for r in range(3, 8):
    print('ROW', r, '->', [ws2.cell(row=r, column=c).value for c in range(1, 8)])
