import openpyxl

path = '/home/vicky/Projects/KRIYA/workspace/CDU-II_C101_Overhead_Vapor_Line_Integrity_Cost_Workbook.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb['Cost & Severity Breakdown']

data_rows = []
for r in range(3, 12):
    tag = ws.cell(row=r, column=1).value
    if tag and str(tag).strip() and str(tag).strip() != 'TOTAL REMEDIATION EXPENDITURE':
        data_rows.append(r)

print('Data rows:', data_rows)
for r in data_rows:
    print('ROW', r, '->', [ws.cell(row=r, column=c).value for c in range(1, 8)])
